"""Generate harmless local preview inputs; never uploads or touches application data."""
import pathlib
import sys
import zipfile

import pymupdf
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches

destination = pathlib.Path(sys.argv[1])
destination.mkdir(parents=True, exist_ok=True)
image = Image.new('RGB', (640, 360), '#e5f4ee')
draw = ImageDraw.Draw(image)
draw.rectangle((35, 35, 605, 325), outline='#15754e', width=8)
draw.text((65, 150), 'LOCAL DOCUMENT PREVIEW - NO UPLOAD', fill='#12442e')
image.save(destination / 'local-image.jpg')

pdf = pymupdf.open()
for number in range(1, 4):
    page = pdf.new_page(width=595, height=842)
    page.insert_text((60, 80), f'LOCAL PDF PREVIEW - PAGE {number}', fontsize=22)
    page.insert_text((60, 125), 'Generated test evidence. This file remains in the browser.', fontsize=11)
    page.draw_rect(pymupdf.Rect(60, 170, 535, 310), color=(0.1, 0.5, 0.3), fill=(0.9, 0.97, 0.94))
    page.insert_text((85, 235), f'Pagination marker {number} of 3', fontsize=18)
pdf.save(destination / 'multipage-preview.pdf')
pdf.save(destination / 'protected-preview.pdf', encryption=pymupdf.PDF_ENCRYPT_AES_256,
         owner_pw='FixtureOwnerOnly', user_pw='PreviewFixtureOnly123')

word = Document()
word.add_heading('Local Word preview', level=1)
paragraph = word.add_paragraph('Formatted preview preserves ')
paragraph.add_run('bold emphasis').bold = True
paragraph.add_run(' and ')
paragraph.add_run('italic emphasis').italic = True
paragraph.add_run(' without uploading the document.')
table = word.add_table(rows=3, cols=2)
table.style = 'Table Grid'
for row, values in zip(table.rows, [('Item', 'Amount'), ('Travel arrangements', '450'), ('Meeting support', '120')]):
    for cell, value in zip(row.cells, values):
        cell.text = value
word.add_picture(str(destination / 'local-image.jpg'), width=Inches(3))
word.add_paragraph('The supporting image is embedded in this Word document.')
word.save(destination / 'formatted-preview.docx')

# A Word relationship that references a remote image must never leak document data.
# It is separate from the normal embedded-image fixture so both paths are covered.
with zipfile.ZipFile(destination / 'formatted-preview.docx') as source:
    with zipfile.ZipFile(destination / 'remote-image-preview.docx', 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            value = source.read(info.filename)
            if info.filename == 'word/_rels/document.xml.rels':
                value = value.replace(b'Target="media/image1.jpg"', b'Target="https://preview-leak.invalid/private-document.jpg" TargetMode="External"')
            target.writestr(info, value)
(destination / 'malformed-preview.pdf').write_bytes(b'%PDF-1.7\nnot a valid PDF document')
(destination / 'malformed-preview.docx').write_bytes(b'PK\x03\x04not a valid Word ZIP archive')
(destination / 'legacy-preview.doc').write_bytes(b'Legacy Word fixture: use download fallback.')
print('LOCAL_PREVIEW_FIXTURES_READY')
